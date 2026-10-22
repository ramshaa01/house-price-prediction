import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Page size: A4 with required university margins:
# Left 3.5 cm (1.38 in), Top 2.5 cm (0.98 in), Right 1.25 cm (0.49 in), Bottom 1.25 cm (0.49 in)
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(3.5 / 2.54)
    section.right_margin = Inches(1.25 / 2.54)
    section.top_margin = Inches(2.5 / 2.54)
    section.bottom_margin = Inches(1.25 / 2.54)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.font.color.rgb = RGBColor(0, 0, 0)
style.paragraph_format.line_spacing = 2.0
style.paragraph_format.space_after = Pt(6)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_chapter(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.font.bold = True
    return p

def add_sec(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    return p

def add_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p

# --- TITLE PAGE ---
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_before = Pt(20)
p_title.paragraph_format.space_after = Pt(24)
r = p_title.add_run("PROJECT REPORT ON\nHOUSE PRICE PREDICTION USING MACHINE\nLEARNING")
r.font.name = 'Times New Roman'
r.font.size = Pt(18)
r.font.bold = True

p_logo = doc.add_paragraph()
p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_logo.paragraph_format.space_after = Pt(28)
r_l = p_logo.add_run("[ AFFIX UNIVERSITY & COLLEGE LOGOS HERE ]")
r_l.font.name = 'Times New Roman'
r_l.font.size = Pt(11)
r_l.font.italic = True
r_l.font.bold = True

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_after = Pt(28)
r_sub = p_sub.add_run("Submitted By:\n\nRamsha Fatima\nRoll No.: 2300101530112\n\n\nUnder the Guidance of:\nMr. Arpit Dubey")
r_sub.font.name = 'Times New Roman'
r_sub.font.size = Pt(13)
r_sub.font.bold = True

p_inst = doc.add_paragraph()
p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_inst.paragraph_format.space_after = Pt(0)
r_inst = p_inst.add_run("Department of Computer Science & Engineering\nUnited College of Engineering and Research, Prayagraj\nAffiliated to Dr. A.P.J. Abdul Kalam Technical University, Lucknow\nIn Association with Innovation & Incubation Hub, MNNIT Foundation\nAcademic Year 2026")
r_inst.font.name = 'Times New Roman'
r_inst.font.size = Pt(12)
r_inst.font.bold = True

doc.add_page_break()

# --- ACKNOWLEDGEMENT ---
p_ack = doc.add_paragraph()
p_ack.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_ack.paragraph_format.space_after = Pt(16)
r_ack = p_ack.add_run("Acknowledgement")
r_ack.font.name = 'Times New Roman'
r_ack.font.size = Pt(16)
r_ack.font.bold = True

add_p(doc, "I would like to express my profound gratitude to everyone who supported me in completing this project.")
add_p(doc, "I am deeply thankful to my guide, Mr. Arpit Dubey, for his continuous encouragement, invaluable insights, and technical guidance throughout the development of this project. His mentorship was instrumental in shaping my understanding of Machine Learning concepts and ensuring the successful execution of this work.")
add_p(doc, "I am also grateful to the Innovation & Incubation Hub Foundation, MNNIT Allahabad, for providing the internship platform and environment for practical technical learning.")
add_p(doc, "Lastly, I extend my heartfelt thanks to my family, friends, and classmates for their unwavering support and motivation during the course of this project.")

p_names = doc.add_paragraph()
p_names.paragraph_format.space_before = Pt(36)
p_names.paragraph_format.line_spacing = 1.3
r_names = p_names.add_run("Ramsha Fatima\nRoll No.: 2300101530112")
r_names.font.name = 'Times New Roman'
r_names.font.size = Pt(12)
r_names.font.bold = True

doc.add_page_break()

# --- ABSTRACT ---
p_ab = doc.add_paragraph()
p_ab.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_ab.paragraph_format.space_after = Pt(16)
r_ab = p_ab.add_run("Abstract")
r_ab.font.name = 'Times New Roman'
r_ab.font.size = Pt(16)
r_ab.font.bold = True

add_p(doc, "The real estate market is highly dynamic and volatile, making the accurate prediction of house prices a challenging yet essential task for buyers, sellers, and investors. This project aims to build a robust Machine Learning system to predict the median house values in various districts using demographic and geographical data. The problem statement focuses on determining how different features such as location (latitude and longitude), median income, total rooms, and proximity to the ocean influence the pricing of a house.")
add_p(doc, "The study utilizes a housing dataset comprising 20,640 records and 10 attributes. The methodology involves an extensive Exploratory Data Analysis (EDA) to understand feature distributions, identify outliers, and detect missing values. A critical preprocessing pipeline was engineered using ColumnTransformer to perform median imputation and standard scaling for numerical features, and frequent value imputation alongside one-hot encoding for categorical variables. Data leakage was strictly prevented by fitting the imputers and scalers only on the training data after an 80-20 train-test split.")
add_p(doc, "To achieve the objective, three Machine Learning algorithms were implemented and compared: Linear Regression (as a baseline model), Random Forest Regressor, and eXtreme Gradient Boosting (XGBoost) Regressor. The models were evaluated using Root Mean Squared Error (RMSE), Mean Absolute Error (MAE), and R Score. The XGBoost model outperformed the others, achieving the highest R score of 0.821 and the lowest RMSE, proving its high reliability in capturing the complex, non-linear relationships in the dataset. The final outcome is a scalable and highly accurate predictive pipeline that can be utilized in real-world real estate applications to estimate property values efficiently.")

doc.add_page_break()

# --- TABLE OF CONTENTS ---
p_toc = doc.add_paragraph()
p_toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_toc.paragraph_format.space_after = Pt(16)
r_toc = p_toc.add_run("Table of Contents")
r_toc.font.name = 'Times New Roman'
r_toc.font.size = Pt(16)
r_toc.font.bold = True

toc_items = [
    ("1. Chapter 1: Introduction", "1"),
    ("   1.1 Background", "1"),
    ("   1.2 Problem Statement", "1"),
    ("   1.3 Need for House Price Prediction", "2"),
    ("   1.4 Objectives", "2"),
    ("   1.5 Scope", "3"),
    ("   1.6 Real-world Applications", "3"),
    ("2. Chapter 2: System Overview", "4"),
    ("   2.1 Complete Workflow of the Project", "4"),
    ("   2.2 Dataset Description", "4"),
    ("   2.3 Features and Target Variable", "5"),
    ("   2.4 Input and Output", "5"),
    ("   2.5 System Architecture", "5"),
    ("3. Chapter 3: Methodology", "6"),
    ("   3.1 Importing Libraries", "6"),
    ("   3.2 Loading the Dataset", "6"),
    ("   3.3 Dataset Exploration and EDA", "6"),
    ("   3.4 Handling Missing Values", "7"),
    ("   3.5 Removing Duplicates and Statistical Summary", "8"),
    ("   3.6 Train-Test Split", "8"),
    ("   3.7 Data Preprocessing Pipeline", "8"),
    ("4. Chapter 4: Machine Learning Models", "10"),
    ("   4.1 Linear Regression (Baseline Model)", "10"),
    ("   4.2 Random Forest Regressor", "10"),
    ("   4.3 XGBoost Regressor", "11"),
    ("5. Chapter 5: Results and Discussion", "12"),
    ("   5.1 Model Evaluation Metrics", "12"),
    ("   5.2 Comparison of Models", "12"),
    ("   5.3 Best-Performing Model and Interpretation", "13"),
    ("6. Chapter 6: Conclusion and Future Scope", "14"),
    ("   6.1 Summary of Work", "14"),
    ("   6.2 Achievements", "14"),
    ("   6.3 Limitations", "14"),
    ("   6.4 Future Improvements", "14"),
    ("7. References", "16"),
    ("8. Appendix", "17")
]

tbl_toc = doc.add_table(rows=len(toc_items), cols=2)
tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
for idx, (title, pg) in enumerate(toc_items):
    c0 = tbl_toc.cell(idx, 0)
    c1 = tbl_toc.cell(idx, 1)
    c0.width = Inches(5.2)
    c1.width = Inches(0.8)
    p0 = c0.paragraphs[0]
    p0.paragraph_format.line_spacing = 1.3
    p0.paragraph_format.space_after = Pt(2)
    r0 = p0.add_run(title)
    r0.font.name = 'Times New Roman'
    r0.font.size = Pt(11)
    if not title.startswith("   "):
        r0.font.bold = True
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.line_spacing = 1.3
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run(pg)
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    if not title.startswith("   "):
        r1.font.bold = True

doc.add_page_break()

# --- CHAPTER 1: INTRODUCTION ---
add_chapter(doc, "Chapter 1: Introduction")
add_sec(doc, "1.1 Background")
add_p(doc, "The real estate industry plays a pivotal role in the global economy. For decades, the valuation of properties was primarily done through manual appraisals, which are time-consuming, expensive, and subject to human bias. With the advent of data science and Machine Learning, predictive modeling has emerged as a powerful tool to automate and optimize property valuation. By analyzing historical housing data, algorithms can identify hidden patterns and trends to forecast property prices accurately.")

add_sec(doc, "1.2 Problem Statement")
add_p(doc, "Housing prices are influenced by a multitude of intersecting factors, ranging from geographical coordinates (latitude and longitude) to socio-economic indicators (median income of the neighborhood). The core problem is to design an automated predictive model that can take these heterogeneous data points as inputs and output a continuous numerical value representing the median house price. The model must handle missing data, categorical text data, and variations in feature scales to ensure accurate predictions.")

add_sec(doc, "1.3 Need for House Price Prediction")
add_p(doc, "• For Buyers: To ensure they are not overpaying for a property.")
add_p(doc, "• For Sellers: To optimally price their property to attract buyers while maximizing profit.")
add_p(doc, "• For Real Estate Agents & Investors: To identify undervalued properties and make data-driven investment decisions.")
add_p(doc, "• For Mortgage Lenders: To accurately assess risk before granting loans backed by property.")

add_sec(doc, "1.4 Objectives")
add_p(doc, "1. To explore and understand the key determinants of house values through Exploratory Data Analysis (EDA).")
add_p(doc, "2. To build an automated data preprocessing pipeline that handles missing values, categorical encoding, and feature scaling without Data Leakage.")
add_p(doc, "3. To train and evaluate multiple regression models: Linear Regression, Random Forest, and XGBoost.")
add_p(doc, "4. To benchmark model performance using standard metrics (RMSE, MAE, R Score) to determine the best-performing model.")

add_sec(doc, "1.5 Scope")
add_p(doc, "The scope of this project is based on the provided California housing dataset. The project covers data cleaning, feature engineering, pipeline creation, model training, and performance evaluation. It demonstrates the complete Machine Learning predictive analysis workflow.")

add_sec(doc, "1.6 Real-world Applications")
add_p(doc, "• Automated Valuation Models (AVMs): Systems like Zillow's Zestimate use similar principles.")
add_p(doc, "• Urban Planning: Helping city planners understand affordability across different districts.")
add_p(doc, "• Economic Forecasting: Tracking real estate trends as a metric for overall economic health.")

doc.add_page_break()

# --- CHAPTER 2: SYSTEM OVERVIEW ---
add_chapter(doc, "Chapter 2: System Overview")
add_sec(doc, "2.1 Complete Workflow of the Project")
add_p(doc, "The project follows a standard Machine Learning lifecycle:")
add_p(doc, "1. Data Ingestion: Loading the raw CSV dataset into a Pandas DataFrame.")
add_p(doc, "2. Data Exploration: Understanding shapes, data types, missing values, and statistical properties.")
add_p(doc, "3. Data Visualization: Plotting distributions of categorical and continuous variables to identify skews and outliers.")
add_p(doc, "4. Data Splitting: Separating the data into training and testing sets before preprocessing to ensure strict model validation.")
add_p(doc, "5. Data Preprocessing: Building an automated pipeline using Scikit-Learn's ColumnTransformer to impute missing values, scale numbers, and encode text.")
add_p(doc, "6. Model Training: Fitting Linear Regression, Random Forest, and XGBoost models to the processed training data.")
add_p(doc, "7. Model Evaluation: Testing the models on unseen data and calculating error metrics.")

add_sec(doc, "2.2 Dataset Description")
add_p(doc, "The dataset used in this project is a historical housing dataset containing information about various block groups (districts) in California.")
add_p(doc, "• Total Records (Rows): 20,640")
add_p(doc, "• Total Attributes (Columns): 10")

add_sec(doc, "2.3 Features and Target Variable")
add_p(doc, "Numerical Features:")
add_p(doc, "• longitude & latitude: Geographical coordinates representing the location of the block.")
add_p(doc, "• housing_median_age: Median age of a house within a block.")
add_p(doc, "• total_rooms: Total number of rooms within a block.")
add_p(doc, "• total_bedrooms: Total number of bedrooms within a block (contains missing values).")
add_p(doc, "• population: Total number of people residing within a block.")
add_p(doc, "• households: Total number of households (a group of people residing in a home unit) for a block.")
add_p(doc, "• median_income: Median income for households within a block of houses (measured in tens of thousands of USD).")
add_p(doc, "Categorical Features:")
add_p(doc, "• ocean_proximity: Location of the house relative to the ocean (e.g., <1H OCEAN, INLAND, NEAR OCEAN, NEAR BAY, ISLAND).")
add_p(doc, "Target Variable:")
add_p(doc, "• median_house_value: The median house value for households within a block (measured in US Dollars). This is the dependent variable our models will attempt to predict.")

add_sec(doc, "2.4 Input and Output")
add_p(doc, "• Input: A vector of 9 features (8 numerical, 1 categorical) representing a specific housing block.")
add_p(doc, "• Output: A single continuous numerical value predicting the median_house_value.")

add_sec(doc, "2.5 System Architecture")
add_p(doc, "The architecture relies heavily on Scikit-Learn's Pipeline object, ensuring that data flows sequentially from raw input to predicted output without manual intervention at each step.")

doc.add_page_break()

# --- CHAPTER 3: METHODOLOGY ---
add_chapter(doc, "Chapter 3: Methodology")
add_p(doc, "This chapter details the exact step-by-step implementation as executed in the project's Jupyter Notebook.")

add_sec(doc, "3.1 Importing Libraries")
add_p(doc, "The project begins by installing and importing essential Python libraries:")
add_p(doc, "• pandas and numpy: For data manipulation, numerical computations, and DataFrame management.")
add_p(doc, "• matplotlib.pyplot and seaborn: For creating visual plots and exploratory charts.")
add_p(doc, "• scikit-learn: Specifically, modules for train_test_split, Pipeline, ColumnTransformer, SimpleImputer, StandardScaler, OneHotEncoder, and regression models.")
add_p(doc, "• xgboost: For utilizing the Extreme Gradient Boosting algorithm.")

add_sec(doc, "3.2 Loading the Dataset")
add_p(doc, "The dataset is loaded directly using Pandas:")
add_p(doc, "df = pd.read_csv('housing.csv')")
add_p(doc, "Initial sanity checks were conducted using df.shape, df.info(), and df.head().")

add_sec(doc, "3.3 Dataset Exploration and EDA")
add_p(doc, "Basic Overview: The initial inspection revealed 20,640 rows and 10 columns. The data types consisted of 9 float64 columns and 1 object column (ocean_proximity).")
add_p(doc, "Data Visualizations:")
add_p(doc, "• sns.countplot was used to plot the frequency of categories in ocean_proximity. The majority of houses were found to be in the <1H OCEAN category, followed by INLAND.")
add_p(doc, "• sns.histplot was used on the target variable (median_house_value). The histogram revealed a right-skewed distribution with a sudden, abnormal peak (higher cap) at $500,001. This indicated that values were capped, a crucial finding because it means models trained on this data might struggle to predict prices beyond $500,000 accurately.")
add_p(doc, "• Correlation Heatmap: Plotted using sns.heatmap on numeric features. The most prominent insight was that median_income had the strongest positive correlation with median_house_value (approx. 0.68), making it the single most informative predictor.")

add_sec(doc, "3.4 Handling Missing Values")
add_p(doc, "Checking Missing Values: Running df.isnull().sum() revealed that total_bedrooms had 207 missing values. All other columns were complete.")
add_p(doc, "How it was handled: Instead of dropping these rows (which would lose 207 rows of valuable data) or manually filling them upfront, the imputation was deferred to the Scikit-Learn Pipeline using SimpleImputer(strategy='median').")
add_p(doc, "Why it was done: Calculating the median on the entire dataset and replacing missing values before splitting causes Data Leakage (where information from the testing set bleeds into the training set). By deferring this to the pipeline, the median is calculated only on the training data and then applied to both sets.")

add_sec(doc, "3.5 Removing Duplicates and Statistical Summary")
add_p(doc, "• Duplicates: df.duplicated().sum() was executed and returned 0, confirming no duplicate entries exist.")
add_p(doc, "• Statistical Summary: df.describe() provided count, mean, standard deviation, minimum, quartiles, and maximum for each numeric feature, revealing varying scales across columns (e.g., housing_median_age ranged from 1 to 52, whereas population ranged up to 35,682).")

add_sec(doc, "3.6 Train-Test Split")
add_p(doc, "The data was split into training and testing sets before applying any scaling or encoding:")
add_p(doc, "• Test Size: 20% (4,128 records)")
add_p(doc, "• Train Size: 80% (16,512 records)")
add_p(doc, "• Random State: 42 (ensuring reproducibility)")

add_sec(doc, "3.7 Data Preprocessing Pipeline")
add_p(doc, "A comprehensive preprocessor was engineered using Scikit-Learn's ColumnTransformer to handle heterogeneous data seamlessly:")
add_p(doc, "Numerical Pipeline: Built for the 8 continuous columns using Pipeline:")
add_p(doc, "1. SimpleImputer(strategy='median'): Fills missing values in total_bedrooms with the training set median.")
add_p(doc, "2. StandardScaler(): Normalizes features to have a mean of 0 and standard deviation of 1. This prevents features with large numbers (like total_rooms) from dominating models over features with small numbers (like median_income).")
add_p(doc, "Categorical Pipeline: Built for the ocean_proximity column:")
add_p(doc, "1. SimpleImputer(strategy='most_frequent'): Acts as a fallback for any missing categorical data by replacing it with the mode.")
add_p(doc, "2. OneHotEncoder(handle_unknown='ignore'): Converts the text-based ocean_proximity into numerical format by creating binary columns for each category (e.g., 1 if INLAND, 0 otherwise).")
add_p(doc, "Column Transformer: These two pipelines were combined into a single preprocess object using ColumnTransformer. This object was fitted strictly on X_train and used to transform both X_train and X_test, ensuring zero data leakage.")

doc.add_page_break()

# --- CHAPTER 4: MACHINE LEARNING MODELS ---
add_chapter(doc, "Chapter 4: Machine Learning Models")
add_sec(doc, "4.1 Linear Regression (Baseline Model)")
add_p(doc, "How it works: Fits a linear equation to the observed data by finding the optimal coefficients for each feature that minimize the sum of squared residuals.")
add_p(doc, "Advantages: Highly interpretable, computationally fast, and provides a clear baseline.")
add_p(doc, "Disadvantages: Assumes a strictly linear relationship between features and the target, making it prone to high bias (underfitting) on complex datasets.")
add_p(doc, "Why it was selected: Serves as the fundamental benchmark against which more sophisticated non-linear models are evaluated.")

add_sec(doc, "4.2 Random Forest Regressor")
add_p(doc, "How it works: An ensemble method that constructs a forest of uncorrelated Decision Trees during training. Each tree is trained on a random bootstrap sample of the data, and feature splits are chosen from a random subset of features. The final prediction is the average of predictions across all 100 individual trees.")
add_p(doc, "Advantages: Reduces overfitting compared to single decision trees, handles non-linear data exceptionally well, and does not require extensive feature scaling.")
add_p(doc, "Disadvantages: Computationally expensive and slower to predict than linear models; acts as a 'black box' with lower interpretability.")
add_p(doc, "Why it was selected: To capture the complex, non-linear geographical and socio-economic relationships in the housing data.")

add_sec(doc, "4.3 XGBoost Regressor (eXtreme Gradient Boosting)")
add_p(doc, "How it works: Unlike Random Forest, where trees are built independently in parallel, XGBoost builds trees sequentially. Each new tree is trained to predict the residual errors made by the combination of all previous trees. It uses gradient descent to minimize a specified loss function.")
add_p(doc, "Configuration in project: Trained with 100 estimators and a maximum depth of 6.")
add_p(doc, "Advantages: Extremely fast due to parallel computing, handles missing values internally, provides built-in regularization to prevent overfitting, and frequently achieves state-of-the-art results on tabular data.")
add_p(doc, "Disadvantages: Requires more memory and careful hyperparameter tuning.")
add_p(doc, "Why it was selected: To push the predictive capability to its maximum potential on the tabular dataset.")

doc.add_page_break()

# --- CHAPTER 5: RESULTS AND DISCUSSION ---
add_chapter(doc, "Chapter 5: Results and Discussion")
add_p(doc, "After training the models via the defined pipelines, their performance was evaluated on both the training and testing datasets.")

add_sec(doc, "5.1 Model Evaluation Metrics")
add_p(doc, "To quantify the performance, three statistical metrics were used:")
add_p(doc, "• Root Mean Squared Error (RMSE): Represents the square root of the variance of the residuals. It heavily penalizes large errors.")
add_p(doc, "• Mean Absolute Error (MAE): The average of the absolute differences between predictions and actual values. It gives a linear idea of the error in dollars.")
add_p(doc, "• R Score (Coefficient of Determination): Represents the proportion of the variance in the dependent variable that is predictable from the independent variables. A score closer to 1.0 indicates a perfect fit.")

add_sec(doc, "5.2 Comparison of Models")
add_p(doc, "A comparative analysis was generated in a Pandas DataFrame at the end of the notebook. The results on the testing data are as follows:")

# Table 5.1
p_t5 = doc.add_paragraph()
p_t5.paragraph_format.line_spacing = 1.0
p_t5.paragraph_format.space_before = Pt(8)
p_t5.paragraph_format.space_after = Pt(4)
r_t5 = p_t5.add_run("Table 5.1: Performance Comparison Across Regression Models")
r_t5.font.name = 'Times New Roman'
r_t5.font.size = Pt(11)
r_t5.font.bold = True

t5 = doc.add_table(rows=4, cols=4)
t5.alignment = WD_TABLE_ALIGNMENT.CENTER
t5_headers = ["Model", "RMSE", "MAE", "R Score"]
t5_data = [
    ["Linear Regression", "~72,688", "~51,000", "~0.612"],
    ["Random Forest", "~49,455", "~32,900", "~0.817"],
    ["XGBoost", "48,496.32", "32,452.53", "0.821"]
]
for c_idx, text in enumerate(t5_headers):
    c = t5.cell(0, c_idx)
    c.text = text
    set_cell_shading(c, "1E3C72")
    set_cell_margins(c, 80, 80, 120, 120)
    for p in c.paragraphs:
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

for r_idx, r_data in enumerate(t5_data, 1):
    for c_idx, val in enumerate(r_data):
        c = t5.cell(r_idx, c_idx)
        c.text = val
        if r_idx == 3:
            set_cell_shading(c, "E6F2FF")
        elif r_idx % 2 == 1:
            set_cell_shading(c, "F5F7FB")
        set_cell_margins(c, 70, 70, 120, 120)
        for p in c.paragraphs:
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(10)
            if r_idx == 3:
                p.runs[0].font.bold = True

add_sec(doc, "5.3 Best-Performing Model and Interpretation")
add_p(doc, "Best Model: The XGBoost Regressor emerged as the definitive winner.")
add_p(doc, "Interpretation:")
add_p(doc, "• The Linear Regression model underfitted the data, yielding an R of only ~0.61. This proves that house prices do not scale in a strictly straight line with features like income or location.")
add_p(doc, "• Random Forest made a massive leap in accuracy (R ~0.817), successfully capturing non-linear relationships.")
add_p(doc, "• XGBoost slightly edged out Random Forest, achieving the lowest MAE ($32,452.53) and the highest R (0.821). This means that on average, the XGBoost model's predictions are off by about $32,452, which is highly acceptable given that house prices span up to $500,001. The gradient boosting mechanism effectively minimized the residual errors left by initial trees.")

doc.add_page_break()

# --- CHAPTER 6: CONCLUSION AND FUTURE SCOPE ---
add_chapter(doc, "Chapter 6: Conclusion and Future Scope")
add_sec(doc, "6.1 Summary of Work")
add_p(doc, "This project successfully implemented an end-to-end Machine Learning pipeline to predict house prices in California. The raw data was visualized, revealing anomalies like price caps. A robust preprocessing pipeline was constructed using Scikit-Learn to impute missing bedrooms and scale numerical figures, preventing data leakage. Three models were trained and rigorously evaluated.")

add_sec(doc, "6.2 Achievements")
add_p(doc, "• Successfully prevented data leakage through proper Pipeline architecture.")
add_p(doc, "• Replaced missing values and encoded categorical geography seamlessly.")
add_p(doc, "• Achieved a high predictive accuracy with an R Score of 0.821 using the XGBoost Regressor, demonstrating the model's strong capability to understand real estate data.")

add_sec(doc, "6.3 Limitations")
add_p(doc, "• Data Capping: The target variable (median_house_value) and housing_median_age were capped at $500,001 and 52 years, respectively. The model learned this artificial ceiling, which may cause it to severely under-predict the value of ultra-luxury homes.")
add_p(doc, "• Static Dataset: The dataset represents a specific historical snapshot and does not account for modern inflation, changing interest rates, or current housing market crashes.")

add_sec(doc, "6.4 Future Improvements")
add_p(doc, "• Outlier Removal: Removing the capped data points ($500,001) during preprocessing could help the model learn the true upper-bound trajectory of house prices.")
add_p(doc, "• Hyperparameter Tuning: Implementing GridSearchCV or RandomizedSearchCV on the XGBoost model to find the mathematically optimal max_depth and learning_rate.")
add_p(doc, "• Geospatial Feature Engineering: Using the latitude and longitude to calculate the distance to major economic hubs (like Silicon Valley or Los Angeles) rather than relying solely on ocean_proximity.")
add_p(doc, "• Deployment: Wrapping the model in a FastAPI or Streamlit backend to serve predictions via a user-friendly web interface.")

doc.add_page_break()

# --- REFERENCES ---
add_chapter(doc, "References")
refs = [
    "1. A. Geron, Hands-On Machine Learning with Scikit-Learn, Keras, and TensorFlow, 2nd ed. Sebastopol, CA, USA: O'Reilly Media, 2019.",
    "2. F. Pedregosa et al., 'Scikit-learn: Machine Learning in Python,' Journal of Machine Learning Research, vol. 12, pp. 2825-2830, 2011.",
    "3. T. Chen and C. Guestrin, 'XGBoost: A Scalable Tree Boosting System,' in Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, San Francisco, CA, USA, 2016, pp. 785-794.",
    "4. L. Breiman, 'Random Forests,' Machine Learning, vol. 45, no. 1, pp. 5-32, 2001.",
    "5. W. McKinney, 'Data Structures for Statistical Computing in Python,' in Proceedings of the 9th Python in Science Conference, Austin, TX, 2010, pp. 51-56.",
    "6. J. D. Hunter, 'Matplotlib: A 2D Graphics Environment,' Computing in Science & Engineering, vol. 9, no. 3, pp. 90-95, 2007.",
    "7. M. Waskom, Seaborn: Statistical Data Visualization, [Online]. Available: https://seaborn.pydata.org/.",
    "8. R. Pace and R. Barry, 'Sparse Spatial Autoregressions,' Statistics & Probability Letters, vol. 33, no. 3, pp. 291-297, 1997."
]
for r in refs:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(r)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)

doc.add_page_break()

# --- APPENDIX ---
add_chapter(doc, "Appendix")
add_sec(doc, "1. Project Summary")
add_p(doc, "This project focuses on predicting California housing prices using demographics and location data. By implementing an automated Scikit-Learn preprocessing pipeline to impute missing values and encode data, the workflow ensures mathematical stability and prevents data leakage. Linear Regression, Random Forest, and XGBoost were compared. XGBoost emerged as the best model, achieving an R score of 0.821 and an MAE of ~32,452, proving gradient boosting's superiority on tabular real-estate data.")

add_sec(doc, "2. Technologies Used")
add_p(doc, "• Programming Language: Python 3.x")
add_p(doc, "• Development Environment: Jupyter Notebook")
add_p(doc, "• Data Manipulation: Pandas, NumPy")
add_p(doc, "• Data Visualization: Matplotlib, Seaborn")
add_p(doc, "• Machine Learning Frameworks: Scikit-Learn, XGBoost")

add_sec(doc, "3. Flowchart of Project Workflow")
add_p(doc, "Raw Dataset (housing.csv) -> Data Exploration & EDA -> Train-Test Split (80/20) -> ColumnTransformer Pipeline (Median Imputer + StandardScaler + OneHotEncoder) -> Model Training (Linear Regression, Random Forest, XGBoost) -> Performance Evaluation (RMSE, MAE, R Score) -> Predicted Median House Value")

out_path = r'C:\Users\ramsh\.gemini\antigravity\scratch\House_Price_Prediction\Documentation\House_Price_Prediction_Report_Exact_Class_Match.docx'
doc.save(out_path)
print('Successfully generated exact class matching Word report at:', out_path)
